import re
import secrets
import uuid
import os
import zipfile
import asyncio
import smtplib
from io import BytesIO
from contextlib import asynccontextmanager
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import openpyxl
from fastapi import Depends, FastAPI, File, HTTPException, Query, Request, UploadFile
from fastapi.responses import FileResponse
from fastapi.responses import Response
import httpx
from argon2.exceptions import VerificationError
from pydantic import BaseModel, Field, field_validator
from sqlalchemy import func, or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, joinedload, selectinload
from starlette.middleware.sessions import SessionMiddleware
from starlette.middleware.trustedhost import TrustedHostMiddleware

from dotenv import load_dotenv

HERE = Path(__file__).parent
load_dotenv(HERE / ".env")

from database import get_db
from models import (
    ChannelSettings,
    Contact,
    Customer,
    Installment,
    Invoice,
    MessageLog,
    MessageTemplate,
    ReminderEvent,
    ReminderPolicy,
    Tenant,
    User,
)
from security import PASSWORDS, credential_cipher, decrypt_secrets, encrypt_secrets, session_secret


LANGUAGES = {"en", "hi", "mr"}
CHANNELS = {"sms", "whatsapp", "email", "telegram"}
CHANNEL_FIELDS = {
    "sms": (
        {"provider", "gateway_url", "sender_id", "test_recipient"},
        {"username", "password", "api_key"},
    ),
    "whatsapp": (
        {"provider", "phone_number_id", "graph_api_version", "template_name", "template_language", "campaign_name", "test_recipient"},
        {"access_token", "api_key"},
    ),
    "email": (
        {"host", "port", "use_ssl", "sender_email", "test_recipient"},
        {"username", "password"},
    ),
    "telegram": ({"chat_id", "test_recipient"}, {"bot_token"}),
}


TEMPLATE_VARIABLES = {"name", "shop", "invoice", "amount", "due_date", "installment_date", "payment_link"}
DEFAULT_TEMPLATES = {
    ("sms", "reminder"): ("", "Hello {name}, invoice {invoice} has {amount} due on {due_date}. Please arrange payment."),
    ("sms", "installment"): ("", "Hello {name}, installment of {amount} for invoice {invoice} is due on {installment_date}."),
    ("whatsapp", "reminder"): ("", "Hello {name}, your balance of {amount} for invoice {invoice} is due on {due_date}."),
    ("whatsapp", "installment"): ("", "Hello {name}, installment of {amount} for invoice {invoice} is due on {installment_date}."),
    ("email", "reminder"): ("Payment reminder", "Hello {name}, invoice {invoice} has {amount} due on {due_date}. Please arrange payment."),
    ("email", "installment"): ("Payment reminder", "Hello {name}, installment of {amount} for invoice {invoice} is due on {installment_date}."),
    ("telegram", "reminder"): ("", "Hello {name}, invoice {invoice} has {amount} due on {due_date}. Please arrange payment."),
    ("telegram", "installment"): ("", "Hello {name}, installment of {amount} for invoice {invoice} is due on {installment_date}."),
}


@asynccontextmanager
async def lifespan(_: FastAPI):
    if os.getenv("APP_ENV") == "production":
        credential_cipher()
    yield


app = FastAPI(title="Dues Desk", version="0.2.0", lifespan=lifespan, docs_url=None, redoc_url=None)
app.add_middleware(
    SessionMiddleware,
    secret_key=session_secret(),
    session_cookie="duesdesk_session",
    same_site="lax",
    https_only=os.getenv("APP_ENV") == "production",
    max_age=60 * 60 * 12,
)
if os.getenv("APP_ENV") == "production":
    allowed_hosts = [host.strip() for host in os.getenv("ALLOWED_HOSTS", "").split(",") if host.strip()]
    if not allowed_hosts:
        raise RuntimeError("ALLOWED_HOSTS must list the public production hostnames.")
    app.add_middleware(TrustedHostMiddleware, allowed_hosts=allowed_hosts)


class SignupRequest(BaseModel):
    tenant_name: str = Field(min_length=2, max_length=160)
    email: str = Field(min_length=5, max_length=320)
    password: str = Field(min_length=12, max_length=128)


class LoginRequest(BaseModel):
    email: str = Field(min_length=5, max_length=320)
    password: str = Field(min_length=1, max_length=128)


class ContactInput(BaseModel):
    name: str | None = Field(default=None, max_length=200)
    phone: str | None = Field(default=None, max_length=40)
    email: str | None = Field(default=None, max_length=320)
    telegram_chat_id: str | None = Field(default=None, max_length=120)
    is_primary: bool = False
    sms_opted_in: bool = False
    whatsapp_opted_in: bool = False
    email_opted_in: bool = False

    @field_validator("email")
    @classmethod
    def valid_contact_email(cls, value: str | None) -> str | None:
        if value and not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", value.strip()):
            raise ValueError("Contact email address is invalid.")
        return value.strip().lower() if value else value

    @field_validator("phone")
    @classmethod
    def valid_contact_phone(cls, value: str | None) -> str | None:
        if value:
            digits = re.sub(r"\D", "", value)
            if not 7 <= len(digits) <= 15:
                raise ValueError("Contact phone must contain between 7 and 15 digits.")
        return value


class InstallmentInput(BaseModel):
    amount: Decimal = Field(gt=0, max_digits=14, decimal_places=2)
    due_date: date
    reminder_days_before: list[int] = Field(default_factory=lambda: [5, 3], max_length=5)


class InvoiceInput(BaseModel):
    invoice_number: str = Field(min_length=1, max_length=120)
    total_amount: Decimal | None = Field(default=None, ge=0, max_digits=14, decimal_places=2)
    paid_amount: Decimal = Field(default=Decimal("0"), ge=0, max_digits=14, decimal_places=2)
    due_amount: Decimal | None = Field(default=None, ge=0, max_digits=14, decimal_places=2)
    due_date: date | None = None
    installments: list[InstallmentInput] = Field(default_factory=list)


class CustomerInput(BaseModel):
    name: str = Field(min_length=1, max_length=200)
    external_id: str | None = Field(default=None, max_length=120)
    company: str | None = Field(default=None, max_length=200)
    location: str | None = Field(default=None, max_length=200)
    language: str = "en"
    reminder_frequency_days: int | None = Field(default=None, ge=1, le=365)
    reminder_dates: list[date] = Field(default_factory=list, max_length=50)
    contacts: list[ContactInput] = Field(default_factory=list, max_length=50)
    invoices: list[InvoiceInput] = Field(default_factory=list, max_length=100)


class CustomerScheduleInput(BaseModel):
    reminder_frequency_days: int | None = Field(default=None, ge=1, le=365)
    reminder_dates: list[date] = Field(default_factory=list, max_length=50)
    reminder_days_before: list[int] | None = Field(default=None, min_length=1, max_length=10)

    @field_validator("reminder_days_before")
    @classmethod
    def valid_reminder_days_before(cls, values: list[int] | None) -> list[int] | None:
        if values is None:
            return None
        if any(day < 1 or day > 30 for day in values):
            raise ValueError("Reminder lead days must be between 1 and 30.")
        return sorted(set(values), reverse=True)


class CustomerRiskInput(BaseModel):
    flagged: bool
    reason: str | None = Field(default=None, max_length=500)


class ChannelInput(BaseModel):
    enabled: bool = False
    provider: str | None = Field(default=None, max_length=40)
    config: dict[str, Any] = Field(default_factory=dict)
    secrets: dict[str, str] = Field(default_factory=dict)
    clear_secrets: bool = False


class ManualMessageInput(BaseModel):
    contact_id: int = Field(gt=0)
    channel: str
    body: str | None = Field(default=None, max_length=5000)
    subject: str | None = Field(default=None, max_length=200)


class LiveSendingInput(BaseModel):
    enabled: bool


class TemplateInput(BaseModel):
    channel: str
    language: str
    stage: str = Field(default="reminder", min_length=1, max_length=40)
    subject: str | None = Field(default=None, max_length=200)
    body: str = Field(min_length=1, max_length=5000)


def current_user(request: Request, db: Session = Depends(get_db)) -> User:
    user_id = request.session.get("user_id")
    if not isinstance(user_id, int):
        raise HTTPException(status_code=401, detail="Sign in to continue.")
    user = db.scalar(select(User).where(User.id == user_id))
    if user is None:
        request.session.clear()
        raise HTTPException(status_code=401, detail="Your session has expired. Sign in again.")
    return user


def require_csrf(request: Request) -> None:
    expected = request.session.get("csrf_token")
    actual = request.headers.get("X-CSRF-Token")
    if not expected or not actual or not secrets.compare_digest(str(expected), actual):
        raise HTTPException(status_code=403, detail="Invalid or missing CSRF token.")


def normalize_email(value: str) -> str:
    normalized = value.strip().lower()
    if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", normalized):
        raise HTTPException(status_code=422, detail="Enter a valid email address.")
    return normalized


def valid_phone(value: str) -> bool:
    return 7 <= len(re.sub(r"\D", "", value)) <= 15


def valid_email(value: str) -> bool:
    return bool(re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", value))


def customer_query(db: Session, tenant_id: int, customer_id: int) -> Customer:
    customer = db.scalar(
        select(Customer)
        .options(selectinload(Customer.contacts), selectinload(Customer.invoices).selectinload(Invoice.installments))
        .where(Customer.id == customer_id, Customer.tenant_id == tenant_id)
    )
    if customer is None:
        raise HTTPException(status_code=404, detail="Customer not found.")
    return customer


def validate_installments(total: Decimal | None, installments: list[InstallmentInput]) -> None:
    if any(day < 1 or day > 30 for item in installments for day in item.reminder_days_before):
        raise HTTPException(status_code=422, detail="Installment reminder lead days must be between 1 and 30.")
    if total is not None and sum((item.amount for item in installments), Decimal("0")) > total:
        raise HTTPException(status_code=422, detail="Installment amounts cannot exceed the invoice total.")


def add_invoice(db: Session, tenant_id: int, customer_id: int, item: InvoiceInput) -> Invoice:
    validate_installments(item.total_amount, item.installments)
    due = item.due_amount
    if due is None and item.total_amount is not None:
        due = max(item.total_amount - item.paid_amount, Decimal("0"))
    invoice = Invoice(
        tenant_id=tenant_id,
        customer_id=customer_id,
        invoice_number=item.invoice_number.strip(),
        total_amount=item.total_amount,
        paid_amount=item.paid_amount,
        due_amount=due,
        due_date=item.due_date,
        status="open" if due is None or due > 0 else "paid",
    )
    db.add(invoice)
    db.flush()
    for number, installment in enumerate(sorted(item.installments, key=lambda value: value.due_date), 1):
        db.add(
            Installment(
                tenant_id=tenant_id,
                invoice_id=invoice.id,
                installment_number=number,
                amount=installment.amount,
                due_date=installment.due_date,
                reminder_days_before=sorted(set(installment.reminder_days_before), reverse=True),
            )
        )
    db.flush()
    return invoice


def serialize_customer(customer: Customer, reminder_days_before: list[int] | None = None) -> dict[str, Any]:
    invoices = []
    for invoice in customer.invoices:
        invoices.append(
            {
                "id": invoice.id,
                "invoice_number": invoice.invoice_number,
                "total_amount": f"{invoice.total_amount:.2f}" if invoice.total_amount is not None else None,
                "paid_amount": f"{invoice.paid_amount:.2f}",
                "due_amount": f"{invoice.due_amount:.2f}" if invoice.due_amount is not None else None,
                "due_date": invoice.due_date.isoformat() if invoice.due_date else None,
                "status": invoice.status,
                "installments": [
                    {
                        "id": row.id,
                        "number": row.installment_number,
                        "amount": f"{row.amount:.2f}",
                        "due_date": row.due_date.isoformat(),
                        "reminder_days_before": row.reminder_days_before,
                        "status": row.status,
                    }
                    for row in sorted(invoice.installments, key=lambda row: row.installment_number)
                ],
            }
        )
    return {
        "id": customer.id,
        "external_id": customer.external_id,
        "name": customer.name,
        "company": customer.company,
        "location": customer.location,
        "language": customer.language,
        "reminder_frequency_days": customer.reminder_frequency_days,
        "reminder_dates": customer.reminder_dates or [],
        "reminder_days_before": reminder_days_before or [10, 7, 5],
        "risk_flagged": customer.risk_flagged,
        "risk_flag_reason": customer.risk_flag_reason,
        "risk_flagged_at": customer.risk_flagged_at.isoformat() if customer.risk_flagged_at else None,
        "contacts": [
            {
                "id": contact.id,
                "name": contact.name,
                "phone": contact.phone,
                "email": contact.email,
                "telegram_chat_id": contact.telegram_chat_id,
                "is_primary": contact.is_primary,
                "sms_opted_in": contact.sms_opted_in,
                "whatsapp_opted_in": contact.whatsapp_opted_in,
                "email_opted_in": contact.email_opted_in,
            }
            for contact in customer.contacts
        ],
        "invoices": invoices,
    }


def validate_template(body: str) -> None:
    names = re.findall(r"\{([^{}]+)\}", body)
    if "{" in re.sub(r"\{[^{}]+\}", "", body) or "}" in re.sub(r"\{[^{}]+\}", "", body):
        raise HTTPException(status_code=422, detail="Template contains an unmatched brace.")
    unknown = sorted(set(names) - TEMPLATE_VARIABLES)
    if unknown:
        raise HTTPException(status_code=422, detail=f"Unsupported template placeholders: {', '.join(unknown)}.")


def parse_amount(value: Any) -> Decimal | None:
    if value is None or not str(value).strip():
        return None
    text = re.sub(r"[^0-9.\-]", "", str(value))
    try:
        parsed = Decimal(text).quantize(Decimal("0.01"))
        return parsed if parsed >= 0 else None
    except (InvalidOperation, ValueError):
        return None


def parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value is None:
        return None
    for pattern in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d %b %Y", "%d-%b-%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(str(value).strip(), pattern).date()
        except ValueError:
            continue
    return None


def header_map(values: tuple[Any, ...]) -> dict[str, int]:
    aliases = {
        "customer": ("customer", "customer name", "dealer", "dealer name", "contact name", "name", "customer/dealer name"),
        "external_id": ("customer id", "cust id", "custid", "dealer id", "customer code", "client id"),
        "company": ("company", "business", "shop", "shop name", "business name"),
        "location": ("location", "city", "town"),
        "phone": ("phone", "mobile", "phone number", "mobile number", "whatsapp"),
        "email": ("email", "email address", "email id"),
        "invoice": ("invoice", "invoice no", "invoice number", "invoice id", "bill no"),
        "total": ("total", "total amount", "invoice amount", "bill amount"),
        "paid": ("paid", "paid amount", "amount paid", "amount received"),
        "due": ("due", "amount due", "due amount", "balance due", "outstanding"),
        "due_date": ("due date", "payment due date"),
        "frequency": ("reminder frequency", "reminder frequency days", "frequency days"),
        "reminder_dates": ("reminder dates", "custom reminder dates"),
        "installment_date": ("installment date", "instalment date", "committed installment date", "committed instalment date"),
        "installment_amount": ("installment amount", "instalment amount", "committed installment amount", "committed instalment amount"),
        "language": ("language", "preferred language"),
        "sms_opted_in": ("sms consent", "sms opted in"),
        "whatsapp_opted_in": ("whatsapp consent", "whatsapp opted in"),
        "email_opted_in": ("email consent", "email opted in"),
        "telegram_chat_id": ("telegram chat id", "telegram id"),
    }
    normalized = [re.sub(r"[^a-z0-9 ]", " ", str(value or "").lower()).strip() for value in values]
    return {
        key: index
        for key, choices in aliases.items()
        if (index := next((i for i, value in enumerate(normalized) if value in choices), None)) is not None
    }


def is_summary_row(row: tuple[Any, ...], columns: dict[str, int]) -> bool:
    summary_labels = {
        "total",
        "grand total",
        "subtotal",
        "sub total",
        "totals",
        "summary",
        "balance carried forward",
        "brought forward",
    }
    identity_values = [
        cell(row, columns, key)
        for key in ("customer", "company", "invoice")
    ]
    for value in identity_values:
        normalized = re.sub(r"[:\s]+$", "", str(value or "").strip().casefold())
        if normalized in summary_labels:
            return True
    return len(header_map(row)) >= 2


def imported_consent(value: Any) -> bool:
    return str(value or "").strip().lower() in {"yes", "true", "1", "y"}


def cell(row: tuple[Any, ...], columns: dict[str, int], field: str) -> Any:
    index = columns.get(field)
    return row[index] if index is not None and index < len(row) else None


def parse_reminder_dates(value: Any) -> list[str]:
    if not value:
        return []
    dates = []
    for entry in re.split(r"[,;\n]", str(value)):
        parsed = parse_date(entry.strip())
        if not parsed:
            raise ValueError(f"Unrecognized reminder date: {entry.strip()}")
        dates.append(parsed.isoformat())
    return sorted(set(dates))


def schedule_reminders(db: Session, tenant_id: int, invoice: Invoice) -> int:
    """Create idempotent scheduled rows; delivery is intentionally a separate worker concern."""
    customer = db.scalar(
        select(Customer).where(Customer.id == invoice.customer_id, Customer.tenant_id == tenant_id)
    )
    if customer is None or invoice.status != "open":
        return 0
    customer_policy = db.scalar(
        select(ReminderPolicy).where(
            ReminderPolicy.tenant_id == tenant_id,
            ReminderPolicy.scope_key == f"customer:{customer.id}",
        )
    )
    policy = customer_policy or db.scalar(
        select(ReminderPolicy).where(
            ReminderPolicy.tenant_id == tenant_id,
            ReminderPolicy.scope_key == "default",
        )
    )
    frequency = customer.reminder_frequency_days or (policy.frequency_days if policy else 3)
    lead_days = policy.default_reminder_days_before if policy and policy.default_reminder_days_before else [10, 7, 5]
    installment_rows = db.scalars(
        select(Installment).where(
            Installment.tenant_id == tenant_id,
            Installment.invoice_id == invoice.id,
            Installment.status == "planned",
        )
    ).all()
    targets: list[tuple[date, int | None, int]] = []
    if installment_rows:
        for installment in installment_rows:
            days_before_values = (
                customer_policy.default_reminder_days_before
                if customer_policy and customer_policy.default_reminder_days_before
                else installment.reminder_days_before
            )
            for days_before in days_before_values:
                targets.append((installment.due_date - timedelta(days=days_before), installment.id, days_before))
    elif customer.reminder_dates:
        targets.extend((date.fromisoformat(day), None, index) for index, day in enumerate(customer.reminder_dates))
    elif invoice.due_date:
        targets.extend((invoice.due_date - timedelta(days=days), None, days) for days in lead_days)
        for number in range(1, (policy.max_reminders if policy else 4) + 1):
            targets.append((invoice.due_date + timedelta(days=frequency * number), None, number))
    channels = db.scalars(
        select(ChannelSettings).where(
            ChannelSettings.tenant_id == tenant_id,
            ChannelSettings.enabled.is_(True),
        )
    ).all()
    created = 0
    if installment_rows:
        db.query(ReminderEvent).filter(
            ReminderEvent.tenant_id == tenant_id,
            ReminderEvent.invoice_id == invoice.id,
            ReminderEvent.installment_id.is_(None),
            ReminderEvent.status == "scheduled",
        ).update({"status": "cancelled"}, synchronize_session=False)
    for channel in channels:
        for scheduled_day, installment_id, ordinal in targets:
            intended_day = scheduled_day
            instant = max(
                datetime.combine(scheduled_day, datetime.min.time(), timezone.utc),
                datetime.now(timezone.utc),
            )
            key = f"{invoice.id}:{installment_id or 0}:{channel.channel}:{intended_day}:{ordinal}"
            exists = db.scalar(
                select(ReminderEvent).where(
                    ReminderEvent.tenant_id == tenant_id,
                    ReminderEvent.idempotency_key == key,
                )
            )
            if exists:
                if exists.status == "cancelled":
                    exists.status = "scheduled"
                    exists.scheduled_at = instant
                    exists.claimed_at = None
                continue
            db.add(
                ReminderEvent(
                    tenant_id=tenant_id,
                    invoice_id=invoice.id,
                    installment_id=installment_id,
                    channel=channel.channel,
                    status="scheduled",
                    scheduled_at=instant,
                    idempotency_key=key,
                )
            )
            created += 1
    return created


def rebuild_invoice_schedule(db: Session, tenant_id: int, invoice: Invoice) -> int:
    db.query(ReminderEvent).filter(
        ReminderEvent.tenant_id == tenant_id,
        ReminderEvent.invoice_id == invoice.id,
        ReminderEvent.status == "scheduled",
    ).update({"status": "cancelled"}, synchronize_session=False)
    db.flush()
    return schedule_reminders(db, tenant_id, invoice)


@app.get("/")
def home() -> FileResponse:
    return FileResponse(HERE / "saas.html")


@app.get("/healthz")
def health() -> dict[str, str]:
    return {"status": "ok"}


@app.get("/sample.xlsx")
def sample_workbook() -> Response:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Dues"
    sheet.append(
        [
            "Customer Name",
            "Customer ID",
            "Company",
            "Location",
            "Phone",
            "Email",
            "SMS Consent",
            "WhatsApp Consent",
            "Email Consent",
            "Telegram Chat ID",
            "Invoice Number",
            "Total Amount",
            "Paid Amount",
            "Due Amount",
            "Due Date",
            "Installment Amount",
            "Installment Date",
            "Reminder Frequency Days",
            "Reminder Dates",
            "Language",
        ]
    )
    sheet.append(
        [
            "Sample Contact",
            "SAMPLE-001",
            "Sample Store",
            "Pune",
            None,
            None,
            "no",
            "no",
            "no",
            None,
            "SAMPLE-INV-001",
            1200,
            300,
            900,
            date.today() + timedelta(days=30),
            300,
            date.today() + timedelta(days=10),
            7,
            None,
            "en",
        ]
    )
    sheet.append(
        [
            "Sample Contact",
            "SAMPLE-001",
            "Sample Store",
            "Pune",
            None,
            None,
            "no",
            "no",
            "no",
            None,
            "SAMPLE-INV-001",
            1200,
            300,
            900,
            date.today() + timedelta(days=30),
            300,
            date.today() + timedelta(days=20),
            7,
            None,
            "en",
        ]
    )
    file = BytesIO()
    book.save(file)
    return Response(
        file.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="dues_desk_import_template.xlsx"'},
    )


@app.get("/auth/csrf")
def csrf_token(request: Request) -> dict[str, str]:
    token = request.session.get("csrf_token")
    if not isinstance(token, str):
        token = secrets.token_urlsafe(32)
        request.session["csrf_token"] = token
    return {"csrf_token": token}


@app.get("/auth/status")
def auth_status(request: Request, db: Session = Depends(get_db)) -> dict[str, Any]:
    user_id = request.session.get("user_id")
    if not isinstance(user_id, int):
        return {"authenticated": False}
    user = db.scalar(select(User).where(User.id == user_id))
    tenant = db.get(Tenant, user.tenant_id) if user is not None else None
    if user is None or tenant is None:
        request.session.clear()
        return {"authenticated": False}
    return {
        "authenticated": True,
        "user": {"id": user.id, "email": user.email, "role": user.role},
        "tenant": {"id": tenant.id, "name": tenant.name, "live_sending_enabled": tenant.live_sending_enabled},
    }


@app.post("/auth/signup", dependencies=[Depends(require_csrf)])
def signup(payload: SignupRequest, request: Request, db: Session = Depends(get_db)) -> dict[str, Any]:
    email = normalize_email(payload.email)
    tenant_name = payload.tenant_name.strip()
    if len(tenant_name) < 2:
        raise HTTPException(status_code=422, detail="Company name must contain at least two characters.")
    tenant = Tenant(name=tenant_name, slug=f"{re.sub(r'[^a-z0-9]+', '-', tenant_name.lower()).strip('-')[:120]}-{uuid.uuid4().hex[:8]}")
    db.add(tenant)
    db.flush()
    user = User(tenant_id=tenant.id, email=email, password_hash=PASSWORDS.hash(payload.password), role="admin")
    db.add(user)
    for (channel, stage), (subject, body) in DEFAULT_TEMPLATES.items():
        db.add(
            MessageTemplate(
                tenant_id=tenant.id,
                channel=channel,
                language="en",
                stage=stage,
                subject=subject or None,
                body=body,
            )
        )
    try:
        db.commit()
    except IntegrityError as error:
        db.rollback()
        raise HTTPException(status_code=409, detail="An account with that email already exists.") from error
    request.session.clear()
    request.session.update({"user_id": user.id, "csrf_token": secrets.token_urlsafe(32)})
    return {"user": {"id": user.id, "email": user.email, "role": user.role}, "tenant": {"id": tenant.id, "name": tenant.name}}


@app.post("/auth/login", dependencies=[Depends(require_csrf)])
def login(payload: LoginRequest, request: Request, db: Session = Depends(get_db)) -> dict[str, Any]:
    email = normalize_email(payload.email)
    user = db.scalar(select(User).where(User.email == email))
    valid = False
    if user is not None:
        try:
            valid = PASSWORDS.verify(user.password_hash, payload.password)
        except (ValueError, TypeError, VerificationError):
            valid = False
    if not valid or user is None:
        raise HTTPException(status_code=401, detail="Email or password is incorrect.")
    request.session.clear()
    request.session.update({"user_id": user.id, "csrf_token": secrets.token_urlsafe(32)})
    return {"user": {"id": user.id, "email": user.email, "role": user.role}}


@app.post("/auth/logout", dependencies=[Depends(require_csrf)])
def logout(request: Request, _: User = Depends(current_user)) -> dict[str, bool]:
    request.session.clear()
    return {"ok": True}


@app.get("/auth/me")
def me(user: User = Depends(current_user), db: Session = Depends(get_db)) -> dict[str, Any]:
    tenant = db.get(Tenant, user.tenant_id)
    return {
        "user": {"id": user.id, "email": user.email, "role": user.role},
        "tenant": {"id": tenant.id, "name": tenant.name, "live_sending_enabled": tenant.live_sending_enabled},
    }


@app.put("/api/settings/live-sending", dependencies=[Depends(require_csrf)])
def set_live_sending(
    payload: LiveSendingInput,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict[str, bool]:
    if user.role != "admin":
        raise HTTPException(status_code=403, detail="Only a workspace administrator can change live sending.")
    tenant = db.get(Tenant, user.tenant_id)
    tenant.live_sending_enabled = payload.enabled
    db.commit()
    return {"live_sending_enabled": tenant.live_sending_enabled}


@app.get("/api/customers")
def list_customers(user: User = Depends(current_user), db: Session = Depends(get_db)) -> list[dict[str, Any]]:
    customers = db.scalars(
        select(Customer)
        .options(selectinload(Customer.contacts), selectinload(Customer.invoices).selectinload(Invoice.installments))
        .where(Customer.tenant_id == user.tenant_id)
        .order_by(Customer.name)
    ).all()
    policies = db.scalars(
        select(ReminderPolicy).where(
            ReminderPolicy.tenant_id == user.tenant_id,
            ReminderPolicy.scope_key != "default",
        )
    ).all()
    default_policy = db.scalar(
        select(ReminderPolicy).where(
            ReminderPolicy.tenant_id == user.tenant_id,
            ReminderPolicy.scope_key == "default",
        )
    )
    default_lead_days = default_policy.default_reminder_days_before if default_policy else [10, 7, 5]
    lead_days_by_customer = {
        row.customer_id: row.default_reminder_days_before
        for row in policies
        if row.customer_id is not None
    }
    return [
        serialize_customer(customer, lead_days_by_customer.get(customer.id, default_lead_days))
        for customer in customers
    ]


@app.post("/api/customers", dependencies=[Depends(require_csrf)])
def create_customer(
    payload: CustomerInput, user: User = Depends(current_user), db: Session = Depends(get_db)
) -> dict[str, Any]:
    if payload.language not in LANGUAGES:
        raise HTTPException(status_code=422, detail="Language must be en, hi, or mr.")
    customer = Customer(
        tenant_id=user.tenant_id,
        external_id=payload.external_id,
        name=payload.name.strip(),
        company=payload.company,
        location=payload.location,
        language=payload.language,
        reminder_frequency_days=payload.reminder_frequency_days,
        reminder_dates=[item.isoformat() for item in payload.reminder_dates],
    )
    db.add(customer)
    db.flush()
    for contact in payload.contacts:
        db.add(Contact(tenant_id=user.tenant_id, customer_id=customer.id, **contact.model_dump()))
    for invoice_data in payload.invoices:
        invoice = add_invoice(db, user.tenant_id, customer.id, invoice_data)
        schedule_reminders(db, user.tenant_id, invoice)
    db.commit()
    default_policy = db.scalar(
        select(ReminderPolicy).where(
            ReminderPolicy.tenant_id == user.tenant_id,
            ReminderPolicy.scope_key == "default",
        )
    )
    return serialize_customer(
        customer_query(db, user.tenant_id, customer.id),
        default_policy.default_reminder_days_before if default_policy else [10, 7, 5],
    )


@app.patch("/api/customers/{customer_id}/schedule", dependencies=[Depends(require_csrf)])
def update_customer_schedule(
    customer_id: int,
    payload: CustomerScheduleInput,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    customer = customer_query(db, user.tenant_id, customer_id)
    customer.reminder_frequency_days = payload.reminder_frequency_days
    customer.reminder_dates = [item.isoformat() for item in payload.reminder_dates]
    if payload.reminder_days_before is not None:
        policy = db.scalar(
            select(ReminderPolicy).where(
                ReminderPolicy.tenant_id == user.tenant_id,
                ReminderPolicy.scope_key == f"customer:{customer.id}",
            )
        )
        if policy is None:
            policy = ReminderPolicy(
                tenant_id=user.tenant_id,
                customer_id=customer.id,
                scope_key=f"customer:{customer.id}",
            )
            db.add(policy)
        policy.default_reminder_days_before = payload.reminder_days_before
    for invoice in customer.invoices:
        rebuild_invoice_schedule(db, user.tenant_id, invoice)
    db.commit()
    policy = db.scalar(
        select(ReminderPolicy).where(
            ReminderPolicy.tenant_id == user.tenant_id,
            ReminderPolicy.scope_key == f"customer:{customer.id}",
        )
    ) or db.scalar(
        select(ReminderPolicy).where(
            ReminderPolicy.tenant_id == user.tenant_id,
            ReminderPolicy.scope_key == "default",
        )
    )
    return serialize_customer(
        customer_query(db, user.tenant_id, customer_id),
        policy.default_reminder_days_before if policy else [10, 7, 5],
    )


@app.patch("/api/customers/{customer_id}/risk", dependencies=[Depends(require_csrf)])
def update_customer_risk(
    customer_id: int,
    payload: CustomerRiskInput,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    customer = customer_query(db, user.tenant_id, customer_id)
    customer.risk_flagged = payload.flagged
    customer.risk_flag_reason = (
        (payload.reason or "").strip() or "Manually flagged by a workspace user."
        if payload.flagged
        else None
    )
    if payload.flagged:
        customer.risk_flagged_at = datetime.now(timezone.utc)
    db.commit()
    return {
        "id": customer.id,
        "risk_flagged": customer.risk_flagged,
        "risk_flag_reason": customer.risk_flag_reason,
        "risk_flagged_at": customer.risk_flagged_at.isoformat() if customer.risk_flagged_at else None,
    }


@app.post("/api/customers/{customer_id}/invoices", dependencies=[Depends(require_csrf)])
def create_invoice(
    customer_id: int,
    payload: InvoiceInput,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    customer = customer_query(db, user.tenant_id, customer_id)
    invoice = add_invoice(db, user.tenant_id, customer.id, payload)
    db.flush()
    schedule_count = schedule_reminders(db, user.tenant_id, invoice)
    db.commit()
    return {"invoice_id": invoice.id, "scheduled_reminders": schedule_count}


@app.post("/api/import", dependencies=[Depends(require_csrf)])
async def import_workbook(
    workbook: UploadFile = File(...),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    if not (workbook.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Upload an .xlsx workbook.")
    try:
        content = await workbook.read(10 * 1024 * 1024 + 1)
        if len(content) > 10 * 1024 * 1024:
            raise HTTPException(status_code=413, detail="Workbook exceeds the 10 MB import limit.")
        book = openpyxl.load_workbook(BytesIO(content), data_only=True, read_only=True)
        rows: list[dict[str, Any]] = []
        for sheet in book.worksheets:
            candidate_rows = list(sheet.iter_rows(min_row=1, max_row=30, values_only=True))
            matching = [
                (header_index, row, header_map(row))
                for header_index, row in enumerate(candidate_rows)
                if "customer" in header_map(row) or "company" in header_map(row)
            ]
            if not matching:
                continue
            header_index, _, columns = max(matching, key=lambda match: len(match[2]))
            for index, row in enumerate(
                sheet.iter_rows(min_row=header_index + 2, values_only=True),
                header_index + 2,
            ):
                if any(value not in (None, "") for value in row) and not is_summary_row(row, columns):
                    rows.append({"row_number": index, "values": row, "columns": columns})
            break
    except (OSError, ValueError, KeyError, zipfile.BadZipFile, openpyxl.utils.exceptions.InvalidFileException) as error:
        raise HTTPException(status_code=400, detail=f"Could not read workbook: {error}") from error
    if not rows:
        raise HTTPException(status_code=422, detail="No customer rows found; include a customer or company column.")

    created_customers = created_invoices = skipped_rows = 0
    existing_customer_ids = set(
        db.scalars(select(Customer.id).where(Customer.tenant_id == user.tenant_id)).all()
    )
    existing_contact_ids = set(
        db.scalars(select(Contact.id).where(Contact.tenant_id == user.tenant_id)).all()
    )
    existing_invoice_ids = set(
        db.scalars(select(Invoice.id).where(Invoice.tenant_id == user.tenant_id)).all()
    )
    updated_customer_ids: set[int] = set()
    updated_contact_ids: set[int] = set()
    updated_invoice_ids: set[int] = set()
    warnings: list[dict[str, Any]] = []
    for item in rows:
        row, columns, row_number = item["values"], item["columns"], item["row_number"]
        customer_name = str(cell(row, columns, "customer") or cell(row, columns, "company") or "").strip()
        company = str(cell(row, columns, "company") or "").strip() or None
        external_id = str(cell(row, columns, "external_id") or "").strip() or None
        phone_value = str(cell(row, columns, "phone") or "").strip()
        email_value = str(cell(row, columns, "email") or "").strip().lower()
        telegram_value = str(cell(row, columns, "telegram_chat_id") or "").strip()
        if phone_value and not valid_phone(phone_value):
            warnings.append({"row": row_number, "message": "Phone number looks invalid and was not saved."})
            phone_value = ""
        if email_value and not valid_email(email_value):
            warnings.append({"row": row_number, "message": "Email address looks invalid and was not saved."})
            email_value = ""
        if not customer_name:
            skipped_rows += 1
            warnings.append({"row": row_number, "message": "Skipped: customer name and company are both missing."})
            continue
        customer = None
        if external_id:
            customer = db.scalar(
                select(Customer).where(
                    Customer.tenant_id == user.tenant_id,
                    func.lower(Customer.external_id) == external_id.casefold(),
                )
            )
        if customer is None:
            customer = db.scalar(
                select(Customer).where(
                    Customer.tenant_id == user.tenant_id,
                    func.lower(func.trim(Customer.name)) == customer_name.casefold(),
                    (
                        func.lower(func.trim(Customer.company)) == company.casefold()
                        if company
                        else Customer.company.is_(None)
                    ),
                )
            )
        if customer is None and (phone_value or email_value or telegram_value):
            contact_filters = []
            if phone_value:
                contact_filters.append(Contact.phone == phone_value)
            if email_value:
                contact_filters.append(func.lower(Contact.email) == email_value)
            if telegram_value:
                contact_filters.append(Contact.telegram_chat_id == telegram_value)
            matching_customer_ids = db.scalars(
                select(Contact.customer_id)
                .where(Contact.tenant_id == user.tenant_id, or_(*contact_filters))
                .distinct()
            ).all()
            if len(matching_customer_ids) == 1:
                customer = db.scalar(
                    select(Customer).where(
                        Customer.id == matching_customer_ids[0],
                        Customer.tenant_id == user.tenant_id,
                    )
                )
        if customer is None:
            language = str(cell(row, columns, "language") or "en").strip().lower()
            customer = Customer(
                tenant_id=user.tenant_id,
                external_id=external_id,
                name=customer_name,
                company=company,
                location=str(cell(row, columns, "location") or "").strip() or None,
                language=language if language in LANGUAGES else "en",
            )
            db.add(customer)
            db.flush()
            created_customers += 1
        else:
            if external_id and not customer.external_id:
                customer.external_id = external_id
            if company:
                customer.company = company
            if cell(row, columns, "location"):
                customer.location = str(cell(row, columns, "location")).strip() or customer.location
            if cell(row, columns, "customer"):
                customer.name = customer_name
            language_value = str(cell(row, columns, "language") or "").strip().lower()
            if language_value in LANGUAGES:
                customer.language = language_value
            if customer.id in existing_customer_ids:
                updated_customer_ids.add(customer.id)

        contact_filters = []
        if phone_value:
            contact_filters.append(Contact.phone == phone_value)
        if email_value:
            contact_filters.append(func.lower(Contact.email) == email_value)
        if telegram_value:
            contact_filters.append(Contact.telegram_chat_id == telegram_value)
        contact = None
        if contact_filters:
            contact_matches = db.scalars(
                select(Contact).where(
                    Contact.tenant_id == user.tenant_id,
                    Contact.customer_id == customer.id,
                    or_(*contact_filters),
                )
            ).all()
            contact = next(
                (existing for existing in contact_matches if phone_value and existing.phone == phone_value),
                None,
            ) or next(
                (existing for existing in contact_matches if email_value and (existing.email or "").lower() == email_value),
                None,
            ) or next(
                (existing for existing in contact_matches if telegram_value and existing.telegram_chat_id == telegram_value),
                None,
            )
        if contact is None and phone_value:
            phone_digits = re.sub(r"\D", "", phone_value)
            existing_contacts = db.scalars(
                select(Contact).where(
                    Contact.tenant_id == user.tenant_id,
                    Contact.customer_id == customer.id,
                )
            ).all()
            normalized_matches = [
                existing
                for existing in existing_contacts
                if existing.phone and re.sub(r"\D", "", existing.phone) == phone_digits
            ]
            if len(normalized_matches) == 1:
                contact = normalized_matches[0]
        if contact is None and (phone_value or email_value or telegram_value):
            contact = Contact(
                tenant_id=user.tenant_id,
                customer_id=customer.id,
                name=customer_name,
                phone=phone_value or None,
                email=email_value or None,
                telegram_chat_id=telegram_value or None,
                is_primary=not customer.contacts,
            )
            db.add(contact)
            db.flush()
        elif contact is not None:
            if customer_name:
                contact.name = customer_name
            if phone_value:
                contact.phone = phone_value
            if email_value:
                contact.email = email_value
            if telegram_value:
                contact.telegram_chat_id = telegram_value
            if contact.id in existing_contact_ids:
                updated_contact_ids.add(contact.id)
        if contact is not None:
            for key, attribute in (
                ("sms_opted_in", "sms_opted_in"),
                ("whatsapp_opted_in", "whatsapp_opted_in"),
                ("email_opted_in", "email_opted_in"),
            ):
                if key in columns:
                    setattr(contact, attribute, imported_consent(cell(row, columns, key)))

        frequency = cell(row, columns, "frequency")
        if frequency not in (None, ""):
            try:
                parsed_frequency = int(frequency)
                if parsed_frequency < 1 or parsed_frequency > 365:
                    raise ValueError
                customer.reminder_frequency_days = parsed_frequency
            except (TypeError, ValueError):
                warnings.append({"row": row_number, "message": "Reminder frequency must be a whole number from 1 to 365 days."})
        try:
            custom_dates = parse_reminder_dates(cell(row, columns, "reminder_dates"))
            if custom_dates:
                customer.reminder_dates = sorted(set((customer.reminder_dates or []) + custom_dates))
        except ValueError as error:
            warnings.append({"row": row_number, "message": str(error)})

        invoice_number = str(cell(row, columns, "invoice") or "").strip()
        total = parse_amount(cell(row, columns, "total"))
        paid = parse_amount(cell(row, columns, "paid"))
        due = parse_amount(cell(row, columns, "due"))
        due_date = parse_date(cell(row, columns, "due_date"))
        if not invoice_number and all(value is None for value in (total, due, due_date)):
            warnings.append({"row": row_number, "message": "Customer/contact imported; no invoice details were present."})
            continue
        if not invoice_number:
            invoice_number = f"IMPORT-{uuid.uuid4().hex[:10].upper()}"
            warnings.append({"row": row_number, "message": "Invoice number missing; created a review-only placeholder."})
        if total is None and due is None:
            warnings.append({"row": row_number, "message": "Invoice has no readable total or amount due; it needs review."})
        if "due_date" in columns and cell(row, columns, "due_date") and due_date is None:
            warnings.append({"row": row_number, "message": "Due date could not be parsed."})
        installment_date = parse_date(cell(row, columns, "installment_date"))
        installment_amount = parse_amount(cell(row, columns, "installment_amount"))
        invoice = db.scalar(
            select(Invoice).where(
                Invoice.tenant_id == user.tenant_id,
                func.lower(func.trim(Invoice.invoice_number)) == invoice_number.casefold(),
            )
        )
        if invoice is not None and invoice.customer_id != customer.id:
            warnings.append({"row": row_number, "message": f"Invoice {invoice_number} already belongs to a different customer in this workspace."})
            continue
        if invoice is None:
            paid_amount = paid if paid is not None else Decimal("0")
            if due is None and total is not None:
                due = max(total - paid_amount, Decimal("0"))
            invoice = Invoice(
                tenant_id=user.tenant_id,
                customer_id=customer.id,
                invoice_number=invoice_number,
                total_amount=total,
                paid_amount=paid_amount,
                due_amount=due,
                due_date=due_date,
                status="needs_review" if total is None and due is None else ("open" if due is None or due > 0 else "paid"),
            )
            db.add(invoice)
            db.flush()
            created_invoices += 1
        else:
            if total is not None:
                invoice.total_amount = total
            if paid is not None:
                invoice.paid_amount = paid
            if due is not None:
                invoice.due_amount = due
            elif total is not None or paid is not None:
                invoice.due_amount = max(
                    (invoice.total_amount or Decimal("0")) - invoice.paid_amount,
                    Decimal("0"),
                )
            if due_date is not None:
                invoice.due_date = due_date
            invoice.status = (
                "needs_review"
                if invoice.total_amount is None and invoice.due_amount is None
                else "open"
                if invoice.due_amount is None or invoice.due_amount > 0
                else "paid"
            )
            if invoice.id in existing_invoice_ids:
                updated_invoice_ids.add(invoice.id)

        if installment_date and installment_amount and installment_amount > 0:
            existing_installments = db.scalars(
                select(Installment).where(
                    Installment.tenant_id == user.tenant_id,
                    Installment.invoice_id == invoice.id,
                )
            ).all()
            if any(row.due_date == installment_date and row.amount == installment_amount for row in existing_installments):
                warnings.append({"row": row_number, "message": "This installment commitment is already recorded."})
            elif invoice.due_amount is not None and sum((row.amount for row in existing_installments), Decimal("0")) + installment_amount > invoice.due_amount:
                warnings.append({"row": row_number, "message": "Installment total would exceed the invoice balance; commitment needs review."})
            else:
                db.add(Installment(
                    tenant_id=user.tenant_id,
                    invoice_id=invoice.id,
                    installment_number=len(existing_installments) + 1,
                    amount=installment_amount,
                    due_date=installment_date,
                    reminder_days_before=[5, 3],
                ))
        elif cell(row, columns, "installment_date") or cell(row, columns, "installment_amount"):
            warnings.append({"row": row_number, "message": "Installment commitment is incomplete; provide both a valid date and positive amount."})
        db.flush()
        rebuild_invoice_schedule(db, user.tenant_id, invoice)
    db.commit()
    return {
        "customers_created": created_customers,
        "customers_updated": len(updated_customer_ids),
        "contacts_updated": len(updated_contact_ids),
        "invoices_created": created_invoices,
        "invoices_updated": len(updated_invoice_ids),
        "rows_skipped": skipped_rows,
        "warnings": warnings,
    }


@app.get("/api/settings/channels")
def get_channel_settings(user: User = Depends(current_user), db: Session = Depends(get_db)) -> list[dict[str, Any]]:
    settings = db.scalars(
        select(ChannelSettings).where(ChannelSettings.tenant_id == user.tenant_id)
    ).all()
    by_channel = {item.channel: item for item in settings}
    return [
        {
            "channel": channel,
            "enabled": by_channel[channel].enabled if channel in by_channel else False,
            "provider": by_channel[channel].provider if channel in by_channel else None,
            "config": by_channel[channel].config if channel in by_channel else {},
            "configured_secrets": {
                key: bool(value)
                for key, value in decrypt_secrets(by_channel[channel].encrypted_secrets).items()
            }
            if channel in by_channel
            else {},
        }
        for channel in sorted(CHANNELS)
    ]


@app.put("/api/settings/channels/{channel}", dependencies=[Depends(require_csrf)])
def save_channel_settings(
    channel: str,
    payload: ChannelInput,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    if channel not in CHANNELS:
        raise HTTPException(status_code=404, detail="Unsupported messaging channel.")
    config_fields, secret_fields = CHANNEL_FIELDS[channel]
    if extra := set(payload.config) - config_fields:
        raise HTTPException(status_code=422, detail=f"Unsupported configuration fields: {', '.join(sorted(extra))}.")
    if extra := set(payload.secrets) - secret_fields:
        raise HTTPException(status_code=422, detail=f"Unsupported credential fields: {', '.join(sorted(extra))}.")
    provider_choices = {
        "sms": {"generic"},
        "whatsapp": {"meta", "aisensy"},
        "email": {"smtp"},
        "telegram": {"telegram"},
    }
    if payload.provider is not None and payload.provider not in provider_choices[channel]:
        raise HTTPException(status_code=422, detail="Unsupported provider for this channel.")
    if any(
        not isinstance(value, (str, int, bool)) or isinstance(value, str) and len(value) > 1000
        for value in payload.config.values()
    ):
        raise HTTPException(status_code=422, detail="Configuration values must be short text, numbers, or booleans.")
    if any(not isinstance(value, str) or len(value) > 4000 for value in payload.secrets.values()):
        raise HTTPException(status_code=422, detail="Credential values must be text no longer than 4000 characters.")
    row = db.scalar(
        select(ChannelSettings).where(
            ChannelSettings.tenant_id == user.tenant_id,
            ChannelSettings.channel == channel,
        )
    )
    if row is None:
        row = ChannelSettings(tenant_id=user.tenant_id, channel=channel)
        db.add(row)
    existing_secrets = {} if payload.clear_secrets else decrypt_secrets(row.encrypted_secrets)
    existing_secrets.update({key: value for key, value in payload.secrets.items() if value})
    row.enabled = payload.enabled
    row.provider = payload.provider
    row.config = payload.config
    row.encrypted_secrets = encrypt_secrets(existing_secrets) if existing_secrets else None
    db.flush()
    if row.enabled:
        invoices = db.scalars(
            select(Invoice).where(Invoice.tenant_id == user.tenant_id, Invoice.status == "open")
        ).all()
        for invoice in invoices:
            rebuild_invoice_schedule(db, user.tenant_id, invoice)
    else:
        db.query(ReminderEvent).filter(
            ReminderEvent.tenant_id == user.tenant_id,
            ReminderEvent.channel == channel,
            ReminderEvent.status == "scheduled",
        ).update({"status": "cancelled"}, synchronize_session=False)
    db.commit()
    return {
        "channel": row.channel,
        "enabled": row.enabled,
        "provider": row.provider,
        "config": row.config,
        "configured_secrets": {key: bool(value) for key, value in existing_secrets.items()},
    }


@app.post("/api/settings/channels/{channel}/test", dependencies=[Depends(require_csrf)])
async def test_channel_settings(
    channel: str,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict[str, str]:
    if user.role != "admin":
        raise HTTPException(status_code=403, detail="Only a workspace administrator can test channel settings.")
    if channel not in CHANNELS:
        raise HTTPException(status_code=404, detail="Unsupported messaging channel.")
    row = db.scalar(
        select(ChannelSettings).where(
            ChannelSettings.tenant_id == user.tenant_id,
            ChannelSettings.channel == channel,
        )
    )
    if row is None or not row.enabled:
        raise HTTPException(status_code=400, detail="Enable and save this channel before sending a test.")
    recipient = str((row.config or {}).get("test_recipient") or "").strip()
    if not recipient:
        raise HTTPException(status_code=400, detail="Set a test recipient before sending a test.")
    credentials = decrypt_secrets(row.encrypted_secrets)
    test_message = "Dues Desk channel configuration test. No customer reminder was sent."
    try:
        from worker import _http_send, _smtp_send

        if channel == "email":
            await asyncio.to_thread(
                _smtp_send,
                row.config or {},
                credentials,
                recipient,
                "Dues Desk configuration test",
                test_message,
            )
        else:
            await _http_send(
                channel,
                row.provider,
                row.config or {},
                credentials,
                recipient,
                test_message,
                ["Dues Desk", "TEST", "configuration test", "0.00", ""],
            )
    except (ValueError, RuntimeError) as error:
        raise HTTPException(status_code=502, detail=f"Channel test failed: {error}") from error
    except (OSError, smtplib.SMTPException, httpx.HTTPError) as error:
        raise HTTPException(status_code=502, detail=f"Channel test failed ({type(error).__name__}).") from error
    return {"status": "sent", "channel": channel, "recipient": recipient}


@app.post("/api/messages/send", dependencies=[Depends(require_csrf)])
async def send_manual_message(
    payload: ManualMessageInput,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict[str, str]:
    if payload.channel not in CHANNELS:
        raise HTTPException(status_code=422, detail="Choose a supported messaging channel.")
    contact = db.scalar(
        select(Contact)
        .options(joinedload(Contact.customer).selectinload(Customer.invoices))
        .where(Contact.id == payload.contact_id, Contact.tenant_id == user.tenant_id)
    )
    if contact is None:
        raise HTTPException(status_code=404, detail="Contact not found.")

    recipient_fields = {
        "sms": (contact.sms_opted_in, contact.phone),
        "whatsapp": (contact.whatsapp_opted_in, contact.phone),
        "email": (contact.email_opted_in, contact.email),
        "telegram": (bool(contact.telegram_chat_id), contact.telegram_chat_id),
    }
    opted_in, recipient = recipient_fields[payload.channel]
    if not opted_in or not recipient:
        raise HTTPException(status_code=400, detail="This contact has no valid destination or applicable recorded consent for the selected channel.")
    if payload.channel in {"sms", "whatsapp"} and not valid_phone(recipient):
        raise HTTPException(status_code=400, detail="The contact does not have a valid phone number for the selected channel.")
    if payload.channel == "email" and not valid_email(recipient):
        raise HTTPException(status_code=400, detail="The contact does not have a valid email address.")

    tenant = db.get(Tenant, user.tenant_id)
    if tenant is None or not tenant.live_sending_enabled:
        raise HTTPException(status_code=400, detail="Enable live customer messaging before sending.")
    setting = db.scalar(
        select(ChannelSettings).where(
            ChannelSettings.tenant_id == user.tenant_id,
            ChannelSettings.channel == payload.channel,
        )
    )
    if setting is None or not setting.enabled:
        raise HTTPException(status_code=400, detail="Enable and save this channel before sending.")

    body = (payload.body or "").strip()
    subject = (payload.subject or "").strip()
    invoice = next(
        (
            item
            for item in sorted(contact.customer.invoices, key=lambda value: (value.due_date or date.max, value.id))
            if item.status == "open" and item.due_amount is not None and item.due_amount > 0
        ),
        None,
    )
    if payload.channel != "whatsapp" and not body:
        template = db.scalar(
            select(MessageTemplate).where(
                MessageTemplate.tenant_id == user.tenant_id,
                MessageTemplate.channel == payload.channel,
                MessageTemplate.language == contact.customer.language,
                MessageTemplate.stage == "reminder",
            )
        )
        if template is None and contact.customer.language != "en":
            template = db.scalar(
                select(MessageTemplate).where(
                    MessageTemplate.tenant_id == user.tenant_id,
                    MessageTemplate.channel == payload.channel,
                    MessageTemplate.language == "en",
                    MessageTemplate.stage == "reminder",
                )
            )
        if template is None:
            raise HTTPException(status_code=400, detail="Create a reminder template for this channel before using a saved template.")
        values = {
            "name": contact.customer.name,
            "shop": contact.customer.company or contact.customer.name,
            "invoice": invoice.invoice_number if invoice else "not recorded",
            "amount": f"{invoice.due_amount:.2f}" if invoice and invoice.due_amount is not None else "not recorded",
            "due_date": invoice.due_date.isoformat() if invoice and invoice.due_date else "not recorded",
            "installment_date": "",
            "payment_link": "",
        }
        try:
            from worker import _render_template

            body = _render_template(template.body, values)
            subject = subject or _render_template(template.subject or "Payment reminder", values)
        except ValueError as error:
            raise HTTPException(status_code=422, detail="The saved reminder template could not be rendered.") from error
    if payload.channel == "whatsapp" and body:
        raise HTTPException(status_code=422, detail="WhatsApp sends use the configured provider-approved template, not free-form text.")
    if payload.channel != "whatsapp" and not body:
        raise HTTPException(status_code=422, detail="Enter a message or leave it blank to use the saved reminder template.")

    logged_body = body
    logged_subject = subject or f"Message from {tenant.name}"
    try:
        from worker import _http_send, _smtp_send

        credentials = decrypt_secrets(setting.encrypted_secrets)
        if payload.channel == "email":
            await asyncio.to_thread(
                _smtp_send,
                setting.config or {},
                credentials,
                recipient,
                logged_subject,
                body,
            )
        elif payload.channel == "whatsapp":
            if invoice is None:
                raise HTTPException(status_code=400, detail="WhatsApp template sends require an open invoice with a positive balance.")
            template_params = [
                contact.customer.name,
                invoice.invoice_number,
                f"payment reminder; due {invoice.due_date.isoformat() if invoice.due_date else 'not recorded'}",
                f"{invoice.due_amount:.2f}",
                "",
            ]
            logged_body = (
                f"Provider-approved template {setting.config.get('template_name') or setting.config.get('campaign_name') or ''}; "
                f"invoice {invoice.invoice_number}; amount {invoice.due_amount:.2f}."
            )
            await _http_send(
                payload.channel,
                setting.provider,
                setting.config or {},
                credentials,
                recipient,
                "",
                template_params,
            )
        else:
            await _http_send(
                payload.channel,
                setting.provider,
                setting.config or {},
                credentials,
                recipient,
                body,
            )
    except HTTPException:
        raise
    except (ValueError, RuntimeError) as error:
        db.add(
            MessageLog(
                tenant_id=user.tenant_id,
                customer_id=contact.customer_id,
                contact_id=contact.id,
                channel=payload.channel,
                direction="manual",
                status="failed",
                recipient=recipient,
                subject=logged_subject if payload.channel == "email" else subject or None,
                body=logged_body or None,
                failure_code=str(error)[:200],
            )
        )
        db.commit()
        raise HTTPException(status_code=502, detail=f"Message send failed: {error}") from error
    except (OSError, smtplib.SMTPException, httpx.HTTPError) as error:
        db.add(
            MessageLog(
                tenant_id=user.tenant_id,
                customer_id=contact.customer_id,
                contact_id=contact.id,
                channel=payload.channel,
                direction="manual",
                status="failed",
                recipient=recipient,
                subject=logged_subject if payload.channel == "email" else subject or None,
                body=logged_body or None,
                failure_code=f"provider_{type(error).__name__.lower()}",
            )
        )
        db.commit()
        raise HTTPException(status_code=502, detail=f"Message send failed ({type(error).__name__}).") from error

    db.add(
        MessageLog(
            tenant_id=user.tenant_id,
            customer_id=contact.customer_id,
            contact_id=contact.id,
            channel=payload.channel,
            direction="manual",
            status="sent",
            recipient=recipient,
            subject=logged_subject if payload.channel == "email" else subject or None,
            body=logged_body or None,
            sent_at=datetime.now(timezone.utc),
        )
    )
    db.commit()
    return {"status": "sent", "channel": payload.channel}


@app.get("/api/messages")
def list_message_history(user: User = Depends(current_user), db: Session = Depends(get_db)) -> list[dict[str, Any]]:
    rows = db.execute(
        select(MessageLog, Customer.name)
        .join(Customer, (Customer.id == MessageLog.customer_id) & (Customer.tenant_id == MessageLog.tenant_id))
        .where(MessageLog.tenant_id == user.tenant_id)
        .order_by(MessageLog.created_at.desc(), MessageLog.id.desc())
        .limit(200)
    ).all()
    return [
        {
            "id": log.id,
            "customer_id": log.customer_id,
            "customer_name": customer_name,
            "channel": log.channel,
            "direction": log.direction,
            "status": log.status,
            "recipient": log.recipient,
            "subject": log.subject,
            "body": log.body,
            "failure_code": log.failure_code,
            "sent_at": log.sent_at.isoformat() if log.sent_at else None,
            "created_at": log.created_at.isoformat(),
        }
        for log, customer_name in rows
    ]


@app.get("/api/templates")
def list_templates(user: User = Depends(current_user), db: Session = Depends(get_db)) -> list[dict[str, Any]]:
    rows = db.scalars(
        select(MessageTemplate)
        .where(MessageTemplate.tenant_id == user.tenant_id)
        .order_by(MessageTemplate.channel, MessageTemplate.language, MessageTemplate.stage)
    ).all()
    return [
        {
            "id": row.id,
            "channel": row.channel,
            "language": row.language,
            "stage": row.stage,
            "subject": row.subject,
            "body": row.body,
        }
        for row in rows
    ]


@app.put("/api/templates", dependencies=[Depends(require_csrf)])
def save_template(
    payload: TemplateInput,
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    if payload.channel not in CHANNELS or payload.language not in LANGUAGES:
        raise HTTPException(status_code=422, detail="Choose a supported channel and language (English, Hindi, or Marathi).")
    if payload.stage not in {"reminder", "installment"}:
        raise HTTPException(status_code=422, detail="Template stage must be reminder or installment.")
    validate_template(payload.body)
    row = db.scalar(
        select(MessageTemplate).where(
            MessageTemplate.tenant_id == user.tenant_id,
            MessageTemplate.channel == payload.channel,
            MessageTemplate.language == payload.language,
            MessageTemplate.stage == payload.stage,
        )
    )
    if row is None:
        row = MessageTemplate(
            tenant_id=user.tenant_id,
            channel=payload.channel,
            language=payload.language,
            stage=payload.stage,
        )
        db.add(row)
    row.subject = payload.subject
    row.body = payload.body
    db.commit()
    return {"id": row.id, "channel": row.channel, "language": row.language, "stage": row.stage}


@app.delete("/api/templates/{template_id}", dependencies=[Depends(require_csrf)])
def delete_template(
    template_id: int, user: User = Depends(current_user), db: Session = Depends(get_db)
) -> dict[str, bool]:
    row = db.scalar(
        select(MessageTemplate).where(
            MessageTemplate.id == template_id,
            MessageTemplate.tenant_id == user.tenant_id,
        )
    )
    if row is None:
        raise HTTPException(status_code=404, detail="Template not found.")
    db.delete(row)
    db.commit()
    return {"ok": True}


@app.put("/api/reminder-policy", dependencies=[Depends(require_csrf)])
def save_reminder_policy(
    frequency_days: int = Query(default=3, ge=1, le=365),
    max_reminders: int = Query(default=4, ge=1, le=30),
    default_reminder_days_before: list[int] = Query(default=[10, 7, 5]),
    user: User = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    if not default_reminder_days_before or any(day < 1 or day > 30 for day in default_reminder_days_before):
        raise HTTPException(status_code=422, detail="Reminder lead days must be between 1 and 30.")
    row = db.scalar(
        select(ReminderPolicy).where(
            ReminderPolicy.tenant_id == user.tenant_id,
            ReminderPolicy.scope_key == "default",
        )
    )
    if row is None:
        row = ReminderPolicy(tenant_id=user.tenant_id, scope_key="default")
        db.add(row)
    row.frequency_days = frequency_days
    row.max_reminders = max_reminders
    row.default_reminder_days_before = sorted(set(default_reminder_days_before), reverse=True)
    invoices = db.scalars(
        select(Invoice).where(Invoice.tenant_id == user.tenant_id, Invoice.status == "open")
    ).all()
    for invoice in invoices:
        rebuild_invoice_schedule(db, user.tenant_id, invoice)
    db.commit()
    return {
        "frequency_days": row.frequency_days,
        "max_reminders": row.max_reminders,
        "default_reminder_days_before": row.default_reminder_days_before,
    }


@app.get("/api/reminder-policy")
def get_reminder_policy(user: User = Depends(current_user), db: Session = Depends(get_db)) -> dict[str, Any]:
    row = db.scalar(
        select(ReminderPolicy).where(
            ReminderPolicy.tenant_id == user.tenant_id,
            ReminderPolicy.scope_key == "default",
        )
    )
    return {
        "frequency_days": row.frequency_days if row else 3,
        "max_reminders": row.max_reminders if row else 4,
        "default_reminder_days_before": row.default_reminder_days_before if row else [10, 7, 5],
    }


@app.get("/api/reminders")
def list_reminders(user: User = Depends(current_user), db: Session = Depends(get_db)) -> list[dict[str, Any]]:
    rows = db.scalars(
        select(ReminderEvent)
        .where(ReminderEvent.tenant_id == user.tenant_id)
        .order_by(ReminderEvent.scheduled_at)
    ).all()
    return [
        {
            "id": row.id,
            "invoice_id": row.invoice_id,
            "installment_id": row.installment_id,
            "channel": row.channel,
            "status": row.status,
            "scheduled_at": row.scheduled_at.isoformat(),
        }
        for row in rows
    ]
