from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path

from sqlalchemy import select

from database import database_url
from saas_app import app
from models import ChannelSettings, Contact, Customer, Installment, Invoice, MessageTemplate, ReminderEvent, Tenant, User


def new_account(client, name, email):
    csrf = client.get("/auth/csrf").json()["csrf_token"]
    response = client.post(
        "/auth/signup",
        headers={"X-CSRF-Token": csrf},
        json={"tenant_name": name, "email": email, "password": "a-secure-passphrase-42"},
    )
    return response, client.get("/auth/csrf").json()["csrf_token"]


def test_default_database_path_does_not_depend_on_shell_working_directory(monkeypatch, tmp_path):
    monkeypatch.delenv("DATABASE_URL", raising=False)
    monkeypatch.chdir(tmp_path)
    expected = f"sqlite:///{(Path(__file__).resolve().parents[1] / 'dues_desk.db').as_posix()}"
    assert database_url() == expected


def test_signup_requires_csrf_and_creates_tenant_admin(client):
    assert client.get("/auth/status").json() == {"authenticated": False}
    response = client.post(
        "/auth/signup",
        json={"tenant_name": "Company", "email": "owner@example.com", "password": "a-secure-passphrase-42"},
    )
    assert response.status_code == 403

    response, csrf = new_account(client, "Company One", "owner@example.com")
    assert response.status_code == 200
    assert response.json()["user"]["role"] == "admin"
    status = client.get("/auth/status").json()
    assert status["authenticated"] is True
    assert status["tenant"]["name"] == "Company One"
    assert client.get("/auth/me").json()["tenant"]["name"] == "Company One"
    assert client.post("/auth/logout", headers={"X-CSRF-Token": csrf}).json() == {"ok": True}
    assert client.get("/api/customers").status_code == 401


def test_login_customer_schedule_policy_and_template_lifecycle(client):
    signup, csrf = new_account(client, "Lifecycle company", "owner@example.com")
    assert signup.status_code == 200
    assert client.post("/auth/logout", headers={"X-CSRF-Token": csrf}).status_code == 200
    csrf = client.get("/auth/csrf").json()["csrf_token"]
    login = client.post(
        "/auth/login",
        headers={"X-CSRF-Token": csrf},
        json={"email": "OWNER@EXAMPLE.COM", "password": "a-secure-passphrase-42"},
    )
    assert login.status_code == 200
    csrf = client.get("/auth/csrf").json()["csrf_token"]
    assert client.post(
        "/auth/login",
        headers={"X-CSRF-Token": csrf},
        json={"email": "owner@example.com", "password": "incorrect-password"},
    ).status_code == 401

    customer = client.post(
        "/api/customers",
        headers={"X-CSRF-Token": csrf},
        json={"name": "Schedule Customer"},
    )
    assert customer.status_code == 200, customer.text
    customer_id = customer.json()["id"]
    assert client.put(
        "/api/settings/channels/sms",
        headers={"X-CSRF-Token": csrf},
        json={
            "enabled": True,
            "provider": "generic",
            "config": {"gateway_url": "https://sms.example.invalid"},
        },
    ).status_code == 200
    due_date = date.today() + timedelta(days=30)
    invoice = client.post(
        f"/api/customers/{customer_id}/invoices",
        headers={"X-CSRF-Token": csrf},
        json={"invoice_number": "SCHEDULE-1", "due_amount": "300", "due_date": due_date.isoformat()},
    )
    assert invoice.status_code == 200, invoice.text
    assert invoice.json()["scheduled_reminders"] == 7
    assert client.get("/api/reminder-policy").json() == {
        "frequency_days": 3,
        "max_reminders": 4,
        "default_reminder_days_before": [10, 7, 5],
    }

    custom_dates = [date.today() + timedelta(days=5), date.today() + timedelta(days=12)]
    updated = client.patch(
        f"/api/customers/{customer_id}/schedule",
        headers={"X-CSRF-Token": csrf},
        json={
            "reminder_frequency_days": 5,
            "reminder_dates": [item.isoformat() for item in custom_dates],
            "reminder_days_before": [10, 5],
        },
    )
    assert updated.status_code == 200, updated.text
    assert updated.json()["reminder_days_before"] == [10, 5]
    active_dates = {
        item["scheduled_at"][:10]
        for item in client.get("/api/reminders").json()
        if item["status"] == "scheduled"
    }
    assert active_dates == {item.isoformat() for item in custom_dates}
    client.patch(
        f"/api/customers/{customer_id}/schedule",
        headers={"X-CSRF-Token": csrf},
        json={"reminder_frequency_days": 5, "reminder_dates": [], "reminder_days_before": [10, 5]},
    )
    active_dates = {
        item["scheduled_at"][:10]
        for item in client.get("/api/reminders").json()
        if item["status"] == "scheduled"
    }
    assert (due_date - timedelta(days=10)).isoformat() in active_dates
    assert (due_date - timedelta(days=5)).isoformat() in active_dates

    policy = client.put(
        "/api/reminder-policy?frequency_days=8&max_reminders=2&default_reminder_days_before=4",
        headers={"X-CSRF-Token": csrf},
    )
    assert policy.status_code == 200
    assert client.get("/api/reminder-policy").json() == {
        "frequency_days": 8,
        "max_reminders": 2,
        "default_reminder_days_before": [4],
    }

    saved = client.put(
        "/api/templates",
        headers={"X-CSRF-Token": csrf},
        json={
            "channel": "email",
            "language": "mr",
            "stage": "reminder",
            "subject": "Invoice {invoice}",
            "body": "Hello {name}, {amount} is due.",
        },
    )
    assert saved.status_code == 200
    initial_templates = client.get("/api/templates").json()
    assert len(initial_templates) == 9
    removed = client.delete(
        f"/api/templates/{saved.json()['id']}",
        headers={"X-CSRF-Token": csrf},
    )
    assert removed.status_code == 200
    assert saved.json()["id"] not in {item["id"] for item in client.get("/api/templates").json()}


def test_sample_workbook_shows_repeatable_installment_columns(client):
    import openpyxl

    response = client.get("/sample.xlsx")
    assert response.status_code == 200
    workbook = openpyxl.load_workbook(BytesIO(response.content), data_only=True)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    assert "Installment Date" in headers
    invoice_column = headers.index("Invoice Number") + 1
    installment_column = headers.index("Installment Date") + 1
    assert sheet.cell(2, invoice_column).value == sheet.cell(3, invoice_column).value
    assert sheet.cell(2, installment_column).value != sheet.cell(3, installment_column).value


def test_customers_and_invoices_are_isolated_by_tenant(client):
    response, csrf = new_account(client, "First Tenant", "first@example.com")
    assert response.status_code == 200
    created = client.post(
        "/api/customers",
        headers={"X-CSRF-Token": csrf},
        json={
            "name": "Asha",
            "contacts": [{"phone": "+919999999999", "is_primary": True}],
            "invoices": [{"invoice_number": "INV-1", "total_amount": "1000", "due_date": (date.today() + timedelta(days=8)).isoformat()}],
        },
    )
    assert created.status_code == 200, created.text
    assert created.json()["invoices"][0]["due_amount"] == "1000.00"

    client.post("/auth/logout", headers={"X-CSRF-Token": csrf})
    response, second_csrf = new_account(client, "Second Tenant", "second@example.com")
    assert response.status_code == 200
    assert client.get("/api/customers").json() == []
    assert client.patch(
        f"/api/customers/{created.json()['id']}/schedule",
        headers={"X-CSRF-Token": second_csrf},
        json={"reminder_frequency_days": 7, "reminder_dates": []},
    ).status_code == 404


def test_channel_secrets_are_encrypted_and_never_returned(client, signed_in):
    client, csrf = signed_in
    secret = "sample-provider-secret-do-not-return"
    response = client.put(
        "/api/settings/channels/telegram",
        headers={"X-CSRF-Token": csrf},
        json={"enabled": True, "config": {"chat_id": "12345"}, "secrets": {"bot_token": secret}},
    )
    assert response.status_code == 200
    assert secret not in response.text

    settings = client.get("/api/settings/channels").json()
    telegram = next(item for item in settings if item["channel"] == "telegram")
    assert telegram["enabled"] is True
    assert telegram["configured_secrets"] == {"bot_token": True}
    assert secret not in str(settings)

    with app.state.test_sessions() as db:
        record = db.scalar(select(ChannelSettings).where(ChannelSettings.channel == "telegram"))
        assert record is not None
        assert b"sample-provider-secret-do-not-return" not in record.encrypted_secrets


def test_templates_validate_placeholders_and_are_tenant_scoped(client, signed_in):
    client, csrf = signed_in
    bad = client.put(
        "/api/templates",
        headers={"X-CSRF-Token": csrf},
        json={"channel": "sms", "language": "en", "body": "Hello {customer_secret}"},
    )
    assert bad.status_code == 422

    saved = client.put(
        "/api/templates",
        headers={"X-CSRF-Token": csrf},
        json={"channel": "sms", "language": "hi", "body": "नमस्ते {name}, बिल {invoice}"},
    )
    assert saved.status_code == 200
    assert any(item["language"] == "hi" for item in client.get("/api/templates").json())


def test_incomplete_excel_import_reports_rows_and_schedules_installment_reminders(client, signed_in):
    client, csrf = signed_in
    setting = client.put(
        "/api/settings/channels/sms",
        headers={"X-CSRF-Token": csrf},
        json={"enabled": True, "config": {"gateway_url": "https://sms.example.invalid/send"}, "secrets": {}},
    )
    assert setting.status_code == 200

    due = date.today() + timedelta(days=30)
    commitment = date.today() + timedelta(days=10)
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["Customer Name", "Customer ID", "Company", "Phone", "Invoice No", "Total Amount", "Due Date", "Installment Amount", "Installment Date", "Reminder Frequency Days", "WhatsApp Consent"])
    sheet.append(["Mira", "M-1", "Mira Stores", "+919800000001", "INV-42", 1200, due, 400, commitment, 6, "yes"])
    second_commitment = commitment + timedelta(days=10)
    sheet.append(["Mira", "M-1", "Mira Stores", "+919800000001", "INV-42", 1200, due, 400, second_commitment, 6, "yes"])
    sheet.append([None, None, None, None, "INV-SKIP", 100, due, None, None, None])
    file_data = BytesIO()
    book.save(file_data)
    file_data.seek(0)

    response = client.post(
        "/api/import",
        headers={"X-CSRF-Token": csrf},
        files={"workbook": ("dues.xlsx", file_data.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["customers_created"] == 1
    assert result["invoices_created"] == 1
    assert result["rows_skipped"] == 1
    assert any("customer name" in item["message"].lower() for item in result["warnings"])

    customer = client.get("/api/customers").json()[0]
    assert customer["reminder_frequency_days"] == 6
    assert len(customer["invoices"][0]["installments"]) == 2
    assert customer["invoices"][0]["installments"][0]["reminder_days_before"] == [5, 3]
    assert customer["contacts"][0]["whatsapp_opted_in"] is True
    reminders = client.get("/api/reminders").json()
    scheduled_dates = {item["scheduled_at"][:10] for item in reminders}
    assert (commitment - timedelta(days=5)).isoformat() in scheduled_dates
    assert (commitment - timedelta(days=3)).isoformat() in scheduled_dates
    assert (second_commitment - timedelta(days=5)).isoformat() in scheduled_dates
    assert (second_commitment - timedelta(days=3)).isoformat() in scheduled_dates
    event_count = len(reminders)
    duplicate = client.post(
        "/api/import",
        headers={"X-CSRF-Token": csrf},
        files={"workbook": ("dues.xlsx", file_data.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert duplicate.status_code == 200, duplicate.text
    assert duplicate.json()["invoices_created"] == 0
    assert len(client.get("/api/reminders").json()) == event_count


def test_excel_reimport_updates_existing_rows_and_skips_summary_and_repeated_headers(client, signed_in):
    import openpyxl

    client, csrf = signed_in
    header = ["Customer Name", "Customer ID", "Phone", "Email", "Invoice Number", "Total Amount", "Paid Amount", "Due Date"]
    due = date.today() + timedelta(days=30)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(header)
    sheet.append(["Jaya", "J-1", "+919800000001", "jaya@example.com", "JAYA-1", 1000, 100, due])
    sheet.append(["Total", None, None, None, "TOTAL", 1000, 100, due])
    sheet.append(header)
    content = BytesIO()
    workbook.save(content)

    first = client.post(
        "/api/import",
        headers={"X-CSRF-Token": csrf},
        files={"workbook": ("dues.xlsx", content.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert first.status_code == 200, first.text
    assert first.json()["customers_created"] == 1
    assert first.json()["invoices_created"] == 1
    assert len(client.get("/api/customers").json()) == 1
    customer = client.get("/api/customers").json()[0]
    assert len(customer["invoices"]) == 1
    assert customer["invoices"][0]["due_amount"] == "900.00"

    updated_due = due + timedelta(days=1)
    update_book = openpyxl.Workbook()
    update_sheet = update_book.active
    update_sheet.append(header)
    update_sheet.append(["Jaya Updated", None, "+919800000002", "jaya@example.com", "jaya-1", 1000, 350, updated_due])
    update_content = BytesIO()
    update_book.save(update_content)
    second = client.post(
        "/api/import",
        headers={"X-CSRF-Token": csrf},
        files={"workbook": ("dues.xlsx", update_content.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert second.status_code == 200, second.text
    assert second.json()["customers_created"] == 0
    assert second.json()["invoices_created"] == 0
    assert second.json()["invoices_updated"] == 1
    customer = client.get("/api/customers").json()[0]
    assert customer["name"] == "Jaya Updated"
    assert len(customer["contacts"]) == 1
    assert customer["contacts"][0]["phone"] == "+919800000002"
    assert len(customer["invoices"]) == 1
    assert customer["invoices"][0]["paid_amount"] == "350.00"
    assert customer["invoices"][0]["due_amount"] == "650.00"
    assert customer["invoices"][0]["due_date"] == updated_due.isoformat()


def test_manual_send_uses_saved_template_and_is_visible_in_history(client, signed_in, monkeypatch):
    client, csrf = signed_in
    created = client.post(
        "/api/customers",
        headers={"X-CSRF-Token": csrf},
        json={
            "name": "Template Customer",
            "contacts": [{"phone": "+919800000003", "sms_opted_in": True}],
            "invoices": [{"invoice_number": "TEMPLATE-1", "due_amount": "125", "due_date": (date.today() + timedelta(days=3)).isoformat()}],
        },
    )
    contact_id = created.json()["contacts"][0]["id"]
    client.put(
        "/api/settings/channels/sms",
        headers={"X-CSRF-Token": csrf},
        json={"enabled": True, "provider": "generic", "config": {"gateway_url": "https://sms.example.invalid"}},
    )
    client.put(
        "/api/settings/live-sending",
        headers={"X-CSRF-Token": csrf},
        json={"enabled": True},
    )
    saved = client.put(
        "/api/templates",
        headers={"X-CSRF-Token": csrf},
        json={"channel": "sms", "language": "en", "stage": "reminder", "body": "Custom {name}: {amount} for {invoice}."},
    )
    assert saved.status_code == 200
    sent = []

    async def record_send(*args):
        sent.append(args)

    monkeypatch.setattr("worker._http_send", record_send)
    response = client.post(
        "/api/messages/send",
        headers={"X-CSRF-Token": csrf},
        json={"contact_id": contact_id, "channel": "sms"},
    )
    assert response.status_code == 200, response.text
    assert sent[0][5] == "Custom Template Customer: 125.00 for TEMPLATE-1."
    history = client.get("/api/messages")
    assert history.status_code == 200
    assert history.json()[0]["status"] == "sent"
    assert history.json()[0]["body"] == sent[0][5]


def test_worker_logs_reminder_and_flags_customer_after_three_overdue_delivery_days(client, signed_in, monkeypatch):
    import asyncio

    import worker

    client, csrf = signed_in
    tenant_id = client.get("/auth/me").json()["tenant"]["id"]
    with app.state.test_sessions() as db:
        tenant = db.get(Tenant, tenant_id)
        tenant.live_sending_enabled = True
        db.add(ChannelSettings(tenant_id=tenant_id, channel="sms", enabled=True, config={"gateway_url": "https://sms.example.invalid"}))
        customer = Customer(tenant_id=tenant_id, name="Overdue Customer", language="en")
        db.add(customer)
        db.flush()
        db.add(Contact(tenant_id=tenant_id, customer_id=customer.id, phone="+919800000004", sms_opted_in=True))
        invoice = Invoice(
            tenant_id=tenant_id,
            customer_id=customer.id,
            invoice_number="OVERDUE-1",
            total_amount="300",
            paid_amount="0",
            due_amount="300",
            due_date=date.today() - timedelta(days=10),
            status="open",
        )
        db.add(invoice)
        db.flush()
        for index, delivered_day in enumerate((date.today() - timedelta(days=2), date.today() - timedelta(days=1))):
            db.add(
                ReminderEvent(
                    tenant_id=tenant_id,
                    invoice_id=invoice.id,
                    channel="sms",
                    status="sent",
                    scheduled_at=datetime.combine(delivered_day, datetime.min.time(), timezone.utc),
                    delivered_at=datetime.combine(delivered_day, datetime.min.time(), timezone.utc),
                    idempotency_key=f"overdue-sent-{index}",
                )
            )
        queued = ReminderEvent(
            tenant_id=tenant_id,
            invoice_id=invoice.id,
            channel="sms",
            scheduled_at=datetime.now(timezone.utc) - timedelta(minutes=1),
            idempotency_key="overdue-final-reminder",
        )
        db.add(queued)
        db.commit()
        customer_id = customer.id

    async def record_delivery(_):
        return None

    monkeypatch.setattr(worker, "_deliver", record_delivery)
    assert asyncio.run(worker.dispatch_once(app.state.test_sessions)) == 1
    customer = client.get("/api/customers").json()[0]
    assert customer["id"] == customer_id
    assert customer["risk_flagged"] is True
    assert "3 delivered payment reminders" in customer["risk_flag_reason"]
    history = client.get("/api/messages").json()
    assert history[0]["direction"] == "reminder"
    assert history[0]["status"] == "sent"
    flagged_at = customer["risk_flagged_at"]
    cleared = client.patch(
        f"/api/customers/{customer_id}/risk",
        headers={"X-CSRF-Token": csrf},
        json={"flagged": False},
    )
    assert cleared.status_code == 200
    assert cleared.json()["risk_flagged"] is False
    assert cleared.json()["risk_flagged_at"] == flagged_at


def test_channel_test_uses_only_explicit_test_recipient(client, signed_in, monkeypatch):
    import worker

    client, csrf = signed_in
    saved = client.put(
        "/api/settings/channels/telegram",
        headers={"X-CSRF-Token": csrf},
        json={
            "enabled": True,
            "config": {"chat_id": "workspace-chat", "test_recipient": "test-chat"},
            "secrets": {"bot_token": "private-token"},
        },
    )
    assert saved.status_code == 200
    calls = []

    async def record_send(*args, **kwargs):
        calls.append((args, kwargs))

    monkeypatch.setattr(worker, "_http_send", record_send)
    response = client.post("/api/settings/channels/telegram/test", headers={"X-CSRF-Token": csrf})
    assert response.status_code == 200, response.text
    assert response.json() == {"status": "sent", "channel": "telegram", "recipient": "test-chat"}
    assert calls[0][0][4] == "test-chat"
    assert "private-token" not in response.text


def test_foreign_tenant_rows_are_not_visible(client, signed_in):
    client, csrf = signed_in
    with app.state.test_sessions() as db:
        original_tenant_id = db.scalar(select(Tenant.id).where(Tenant.name == "Northwind Traders"))
    assert original_tenant_id is not None

    client.post("/auth/logout", headers={"X-CSRF-Token": csrf})
    response, second_csrf = new_account(client, "Second company", "other@example.com")
    assert response.status_code == 200
    assert client.get("/api/customers").json() == []
    with app.state.test_sessions() as db:
        customer = Customer(tenant_id=original_tenant_id, name="Tenant one's record")
        db.add(customer)
        db.commit()
        customer_id = customer.id
    assert client.patch(
        f"/api/customers/{customer_id}/schedule",
        headers={"X-CSRF-Token": second_csrf},
        json={"reminder_frequency_days": 4, "reminder_dates": []},
    ).status_code == 404


def test_worker_requires_live_mode_and_contact_consent(client, signed_in, monkeypatch):
    import asyncio

    import worker

    client, _ = signed_in
    tenant_id = client.get("/auth/me").json()["tenant"]["id"]
    with app.state.test_sessions() as db:
        db.add(ChannelSettings(tenant_id=tenant_id, channel="sms", enabled=True, config={"gateway_url": "https://sms.example.invalid"}))
        customer = Customer(tenant_id=tenant_id, name="Consent check", language="en")
        db.add(customer)
        db.flush()
        db.add(Contact(tenant_id=tenant_id, customer_id=customer.id, phone="+919999999999", sms_opted_in=False))
        invoice = Invoice(
            tenant_id=tenant_id,
            customer_id=customer.id,
            invoice_number="CONSENT-1",
            due_amount="250.00",
            total_amount="250.00",
            paid_amount="0",
            due_date=date.today(),
            status="open",
        )
        db.add(invoice)
        db.flush()
        customer_id, invoice_id = customer.id, invoice.id
        event = ReminderEvent(
            tenant_id=tenant_id,
            invoice_id=invoice.id,
            channel="sms",
            scheduled_at=datetime.now(timezone.utc) - timedelta(minutes=1),
            idempotency_key="consent-check-event",
        )
        db.add(event)
        db.commit()
        event_id = event.id

    calls = []

    async def record_delivery(item):
        calls.append(item)

    monkeypatch.setattr(worker, "_deliver", record_delivery)
    asyncio.run(worker.dispatch_once(app.state.test_sessions))
    with app.state.test_sessions() as db:
        saved = db.get(ReminderEvent, event_id)
        assert saved.status == "failed"
        assert saved.failure_code == "live_sending_disabled"
        tenant = db.get(Tenant, tenant_id)
        tenant.live_sending_enabled = True
        contact = db.scalar(select(Contact).where(Contact.customer_id == customer_id))
        contact.sms_opted_in = True
        template = db.scalar(
            select(MessageTemplate).where(
                MessageTemplate.tenant_id == tenant_id,
                MessageTemplate.channel == "sms",
                MessageTemplate.language == "en",
                MessageTemplate.stage == "reminder",
            )
        )
        template.body = "Hello {name}, {invoice} has {amount} due on {due_date}."
        next_event = ReminderEvent(
            tenant_id=tenant_id,
            invoice_id=invoice_id,
            channel="sms",
            scheduled_at=datetime.now(timezone.utc) - timedelta(minutes=1),
            idempotency_key="consent-check-event-2",
        )
        db.add(next_event)
        db.commit()
        next_event_id = next_event.id

    asyncio.run(worker.dispatch_once(app.state.test_sessions))
    assert len(calls) == 1
    assert calls[0]["recipient"] == "+919999999999"
    assert "Consent check" in calls[0]["message"]
    with app.state.test_sessions() as db:
        assert db.get(ReminderEvent, next_event_id).status == "sent"
    assert asyncio.run(worker.dispatch_once(app.state.test_sessions)) == 0


def test_manual_messages_send_through_all_configured_channels(client, signed_in, monkeypatch):
    client, csrf = signed_in
    customer = client.post(
        "/api/customers",
        headers={"X-CSRF-Token": csrf},
        json={
            "name": "Manual Send Customer",
            "contacts": [
                {
                    "phone": "+919999999999",
                    "email": "customer@example.com",
                    "telegram_chat_id": "chat-123",
                    "is_primary": True,
                    "sms_opted_in": True,
                    "whatsapp_opted_in": True,
                    "email_opted_in": True,
                }
            ],
            "invoices": [
                {
                    "invoice_number": "MANUAL-1",
                    "total_amount": "250.00",
                    "due_amount": "250.00",
                    "due_date": (date.today() + timedelta(days=5)).isoformat(),
                }
            ],
        },
    )
    assert customer.status_code == 200, customer.text
    contact_id = customer.json()["contacts"][0]["id"]

    channel_settings = {
        "sms": ("generic", {"gateway_url": "https://sms.example.invalid"}, {"api_key": "sms-key"}),
        "whatsapp": (
            "meta",
            {
                "phone_number_id": "123456",
                "graph_api_version": "v21.0",
                "template_name": "payment_reminder",
                "template_language": "en",
            },
            {"access_token": "whatsapp-token"},
        ),
        "email": (
            "smtp",
            {"host": "smtp.example.invalid", "port": 465, "sender_email": "billing@example.com"},
            {"username": "billing@example.com", "password": "mail-password"},
        ),
        "telegram": ("telegram", {}, {"bot_token": "telegram-token"}),
    }
    for channel, (provider, config, secrets) in channel_settings.items():
        response = client.put(
            f"/api/settings/channels/{channel}",
            headers={"X-CSRF-Token": csrf},
            json={"enabled": True, "provider": provider, "config": config, "secrets": secrets},
        )
        assert response.status_code == 200, response.text
    assert client.put(
        "/api/settings/live-sending",
        headers={"X-CSRF-Token": csrf},
        json={"enabled": True},
    ).status_code == 200

    calls = []

    async def record_http(*args):
        calls.append(("http", args))

    def record_smtp(*args):
        calls.append(("smtp", args))

    monkeypatch.setattr("worker._http_send", record_http)
    monkeypatch.setattr("worker._smtp_send", record_smtp)
    messages = {
        "sms": {"body": "Your balance is due."},
        "whatsapp": {"body": None},
        "email": {"subject": "Invoice follow-up", "body": "Your balance is due."},
        "telegram": {"body": "Your balance is due."},
    }
    for channel, message in messages.items():
        response = client.post(
            "/api/messages/send",
            headers={"X-CSRF-Token": csrf},
            json={"contact_id": contact_id, "channel": channel, **message},
        )
        assert response.status_code == 200, response.text
        assert response.json() == {"status": "sent", "channel": channel}
        assert all(
            secret not in response.text
            for secret in ("sms-key", "whatsapp-token", "mail-password", "telegram-token")
        )

    http_calls = [args for kind, args in calls if kind == "http"]
    assert [args[0] for args in http_calls] == ["sms", "whatsapp", "telegram"]
    assert http_calls[0][4:6] == ("+919999999999", "Your balance is due.")
    assert http_calls[1][4] == "+919999999999"
    assert http_calls[1][6] == [
        "Manual Send Customer",
        "MANUAL-1",
        f"payment reminder; due {(date.today() + timedelta(days=5)).isoformat()}",
        "250.00",
        "",
    ]
    assert http_calls[2][4:6] == ("chat-123", "Your balance is due.")
    smtp_call = next(args for kind, args in calls if kind == "smtp")
    assert smtp_call[2:] == ("customer@example.com", "Invoice follow-up", "Your balance is due.")


def test_manual_send_requires_consent_live_mode_and_enabled_channel(client, signed_in):
    client, csrf = signed_in
    created = client.post(
        "/api/customers",
        headers={"X-CSRF-Token": csrf},
        json={"name": "Manual Send Guard", "contacts": [{"phone": "+919999999999"}]},
    )
    contact_id = created.json()["contacts"][0]["id"]
    payload = {"contact_id": contact_id, "channel": "sms", "body": "Please review your balance."}

    response = client.post("/api/messages/send", headers={"X-CSRF-Token": csrf}, json=payload)
    assert response.status_code == 400
    assert "consent" in response.json()["detail"].lower()

    with app.state.test_sessions() as db:
        contact = db.get(Contact, contact_id)
        contact.sms_opted_in = True
        db.commit()
    response = client.post("/api/messages/send", headers={"X-CSRF-Token": csrf}, json=payload)
    assert response.status_code == 400
    assert "live customer messaging" in response.json()["detail"].lower()

    assert client.put(
        "/api/settings/live-sending",
        headers={"X-CSRF-Token": csrf},
        json={"enabled": True},
    ).status_code == 200
    response = client.post("/api/messages/send", headers={"X-CSRF-Token": csrf}, json=payload)
    assert response.status_code == 400
    assert "enable and save" in response.json()["detail"].lower()


def test_manual_send_surfaces_provider_errors(client, signed_in, monkeypatch):
    client, csrf = signed_in
    created = client.post(
        "/api/customers",
        headers={"X-CSRF-Token": csrf},
        json={
            "name": "Provider Error Customer",
            "contacts": [{"phone": "+919999999999", "sms_opted_in": True}],
        },
    )
    contact_id = created.json()["contacts"][0]["id"]
    client.put(
        "/api/settings/channels/sms",
        headers={"X-CSRF-Token": csrf},
        json={"enabled": True, "provider": "generic", "config": {}, "secrets": {}},
    )
    client.put(
        "/api/settings/live-sending",
        headers={"X-CSRF-Token": csrf},
        json={"enabled": True},
    )

    async def fail_send(*_):
        raise ValueError("provider_rejected_message")

    monkeypatch.setattr("worker._http_send", fail_send)
    response = client.post(
        "/api/messages/send",
        headers={"X-CSRF-Token": csrf},
        json={"contact_id": contact_id, "channel": "sms", "body": "Hello."},
    )
    assert response.status_code == 502
    assert "provider_rejected_message" in response.json()["detail"]
    history = client.get("/api/messages")
    assert history.status_code == 200
    assert history.json()[0]["status"] == "failed"
    assert history.json()[0]["failure_code"] == "provider_rejected_message"


def test_sms_and_meta_requests_use_configured_authorization_credentials(monkeypatch):
    import asyncio

    import worker

    requests = []

    class Response:
        status_code = 200

    class Client:
        async def __aenter__(self):
            return self

        async def __aexit__(self, *_):
            return None

        async def post(self, url, **kwargs):
            requests.append((url, kwargs))
            return Response()

    monkeypatch.setattr(worker.httpx, "AsyncClient", lambda **_: Client())
    monkeypatch.setattr(worker, "_validate_public_host", lambda *_: None)
    asyncio.run(
        worker._http_send(
            "sms",
            "generic",
            {"gateway_url": "https://sms.example.invalid"},
            {"api_key": "test-sms-api-key"},
            "+919999999999",
            "SMS test",
        )
    )
    asyncio.run(
        worker._http_send(
            "whatsapp",
            "meta",
            {
                "phone_number_id": "123456",
                "graph_api_version": "v21.0",
                "template_name": "payment_reminder",
                "template_language": "en",
            },
            {"access_token": "test-meta-access-token"},
            "+919999999999",
            "",
            ["Customer", "INV-1", "payment reminder", "100.00", ""],
        )
    )

    assert requests[0][1]["headers"]["Authorization"] == "Bearer test-sms-api-key"
    assert requests[1][1]["headers"]["Authorization"] == "Bearer test-meta-access-token"
